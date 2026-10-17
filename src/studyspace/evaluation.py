"""Evaluation harness (member 6's area): extraction against labelled tasks, planner against reference cases.

Extraction is scored per gold task:
- found        a candidate from the same message(s) describes it
- right day    the candidate's day equals the labelled day (both empty also counts)
- right time   the exact time matches, or both are empty because the time is not stated or ambiguous
- merged       the candidate combines exactly the labelled messages (duplicates, deadline changes)
plus precision (how many candidates match a labelled task) and every failure listed for the report.

Synthetic results are for development only. The figures in Evaluation.pdf come from the held-out real set.
"""

from __future__ import annotations

import json
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from .extraction.merge import title_tokens
from .models import Candidate, PlanRequest
from .timeutil import local

# member 3 labels in Chinese; these words let a label be matched to an English candidate title
_ZH_HINTS = {
    "期中": "midterm", "考试": "exam", "模拟": "mock", "下载": "download paper", "试卷": "paper", "小抄": "cheat sheet",
    "姓名": "write name", "学号": "write", "问卷": "survey", "作业": "assignment", "提交": "submit", "报告": "report",
    "小测": "quiz", "阅读": "read reading", "汇报": "presentation", "项目": "project",
}


@dataclass
class GoldTask:
    gold_id: str
    source_ids: list[str]
    title: str
    due_date: date | None
    due_time: str | None  # "HH:MM"
    kind: str = "deliverable"
    flags: list[str] = field(default_factory=list)
    course_code: str | None = None
    note: str = ""


def load_gold_json(path: str | Path) -> tuple[list[GoldTask], dict[str, str]]:
    raw = json.loads(Path(path).read_text(encoding="utf-8"))
    tasks = [GoldTask(gold_id=t["gold_id"], source_ids=t["source_ids"], title=t["title"],
                      due_date=date.fromisoformat(t["due_date"]) if t.get("due_date") else None, due_time=t.get("due_time"),
                      kind=t.get("kind", "deliverable"), flags=t.get("flags", []), course_code=t.get("course_code"),
                      note=t.get("note", "")) for t in raw["tasks"]]
    return tasks, raw.get("non_task_sources", {})


def load_member3_xlsx(path: str | Path, sheet: str = "标注对照_3号填写") -> list[GoldTask]:
    """Member 3's label workbook (columns task_id, record_ids, course_code, task_name, deadline_date, deadline_time...)."""
    from openpyxl import load_workbook  # optional dependency, only needed for the real set

    ws = load_workbook(path, data_only=True)[sheet]
    rows = list(ws.iter_rows(values_only=True))
    head = [str(h).strip() if h else "" for h in rows[0]]
    out = []
    for row in rows[1:]:
        r = dict(zip(head, row))
        if not r.get("task_id"):
            continue
        d, t = r.get("deadline_date"), r.get("deadline_time")
        due_date = d.date() if isinstance(d, datetime) else (date.fromisoformat(str(d)[:10]) if d else None)
        due_time = t.strftime("%H:%M") if hasattr(t, "strftime") else (str(t)[:5] if t else None)
        status = str(r.get("status") or "")
        flags = ["MIDNIGHT_AMBIGUITY"] if "midnight" in status.lower() else []
        name = str(r.get("task_name") or "")
        out.append(GoldTask(gold_id=str(r["task_id"]), source_ids=[s.strip() for s in str(r.get("record_ids") or "").split(",") if s.strip()],
                            title=name, due_date=due_date, due_time=due_time, flags=flags, course_code=r.get("course_code"),
                            kind="event" if name.startswith("参加") else "deliverable", note=str(r.get("relation_or_event") or "")))
    return out


def _tokens(title: str) -> set[str]:
    extra = " ".join(v for k, v in _ZH_HINTS.items() if k in title)
    return title_tokens(f"{title} {extra}")


def _similarity(a: str, b: str) -> float:
    ta, tb = _tokens(a), _tokens(b)
    return len(ta & tb) / len(ta | tb) if ta and tb else 0.0


def _time(c: Candidate) -> str | None:
    return local(c.due_at).strftime("%H:%M") if c.due_at else None


def evaluate_extraction(candidates: list[Candidate], gold: list[GoldTask], non_task: dict[str, str] | None = None) -> dict:
    non_task = non_task or {}
    pairs = []
    for g in gold:
        for c in candidates:
            shared = set(g.source_ids) & set(c.source_ids)
            if not shared:
                continue
            score = (3 * _similarity(g.title, c.title) + (1.5 if c.due_date == g.due_date else 0)
                     + (0.5 if _time(c) == g.due_time else 0) + (0.5 if c.kind == g.kind else 0)
                     + len(shared) / max(len(g.source_ids), 1))
            pairs.append((score, g.gold_id, c.candidate_id))
    matched: dict[str, Candidate] = {}
    used: set[str] = set()
    by_id = {c.candidate_id: c for c in candidates}
    for score, gid, cid in sorted(pairs, reverse=True):
        if gid in matched or cid in used or score < 1.2:
            continue
        matched[gid] = by_id[cid]
        used.add(cid)

    failures = []
    counts = {"found": 0, "right_day": 0, "right_time": 0, "merged": 0, "flags": 0}
    for g in gold:
        c = matched.get(g.gold_id)
        if c is None:
            failures.append({"type": "missed", "gold_id": g.gold_id, "expected": g.title, "got": None, "note": g.note})
            continue
        counts["found"] += 1
        day_ok = c.due_date == g.due_date
        time_ok = _time(c) == g.due_time if day_ok else False
        merged_ok = set(c.source_ids) == set(g.source_ids)
        flags_ok = set(g.flags) <= set(c.flags)
        counts["right_day"] += day_ok
        counts["right_time"] += time_ok
        counts["merged"] += merged_ok
        counts["flags"] += flags_ok
        base = {"gold_id": g.gold_id, "expected": g.title, "got": c.title, "evidence": c.evidence, "note": g.note}
        if not day_ok:
            failures.append({**base, "type": "wrong_day", "expected_value": str(g.due_date), "got_value": str(c.due_date)})
        elif not time_ok:
            failures.append({**base, "type": "wrong_time", "expected_value": g.due_time, "got_value": _time(c)})
        if not merged_ok:
            failures.append({**base, "type": "wrong_merge", "expected_value": g.source_ids, "got_value": c.source_ids})
        if not flags_ok:
            failures.append({**base, "type": "missing_flag", "expected_value": g.flags, "got_value": c.flags})
    extras = [c for c in candidates if c.candidate_id not in used]
    for c in extras:
        kind = "non_task_flagged" if set(c.source_ids) <= set(non_task) else "extra"
        failures.append({"type": kind, "gold_id": None, "expected": None, "got": c.title, "evidence": c.evidence,
                         "note": "; ".join(non_task.get(s, "") for s in c.source_ids if s in non_task)})
    n = len(gold) or 1
    return {
        "gold_tasks": len(gold),
        "candidates": len(candidates),
        "recall": counts["found"] / n,
        "precision": len(used) / (len(candidates) or 1),
        "right_day": counts["right_day"] / n,
        "right_day_and_time": counts["right_time"] / n,
        "merge_correct": counts["merged"] / n,
        "flags_correct": counts["flags"] / n,
        "counts": counts,
        "failures": failures,
    }


def evaluate_planner(cases_path: str | Path) -> dict:
    """Run the engine on member 4's reference cases and check each plan independently."""
    from .planner import check_plan, generate

    raw = json.loads(Path(cases_path).read_text(encoding="utf-8"))
    rows = []
    for case in raw["cases"]:
        request = PlanRequest.model_validate({**case["input"], "settings": raw["settings"]})
        result = generate(request)
        issues = check_plan(request, result)
        scheduled = sum((b.end - b.start).total_seconds() / 60 for b in result.blocks)
        missing = sum(u.minutes for u in result.unscheduled)
        rows.append({
            "case": case["id"], "status": result.status, "expected_status": case["reference_plan"]["status"],
            "scheduled": scheduled, "expected_scheduled": case["expected"]["scheduled_minutes"],
            "unscheduled": missing, "expected_unscheduled": case["expected"]["unscheduled_minutes"],
            "rule_issues": issues,
            "passed": not issues and result.status == case["reference_plan"]["status"]
            and scheduled == case["expected"]["scheduled_minutes"] and missing == case["expected"]["unscheduled_minutes"],
        })
    return {"cases": rows, "passed": sum(r["passed"] for r in rows), "total": len(rows)}


def markdown_report(title: str, extraction: dict | None, planner: dict | None) -> str:
    lines = [f"# {title}", ""]
    if extraction:
        e = extraction
        lines += [
            "## Extraction", "",
            "| Measure | Value |", "| --- | --- |",
            f"| Labelled tasks | {e['gold_tasks']} |",
            f"| Candidates produced | {e['candidates']} |",
            f"| Found (recall) | {e['recall']:.0%} ({e['counts']['found']}/{e['gold_tasks']}) |",
            f"| Precision | {e['precision']:.0%} |",
            f"| Right day | {e['right_day']:.0%} |",
            f"| Right day and time | {e['right_day_and_time']:.0%} |",
            f"| Messages merged correctly | {e['merge_correct']:.0%} |",
            f"| Warning flags present | {e['flags_correct']:.0%} |",
            "", f"### Failure cases ({len(e['failures'])})", "",
            "| # | Type | Label | System output | Expected | Got | Evidence |", "| --- | --- | --- | --- | --- | --- | --- |",
        ]
        for i, f in enumerate(e["failures"], 1):
            ev = re.sub(r"\s+", " ", f.get("evidence") or "")[:90].replace("|", "/")
            lines.append(f"| {i} | {f['type']} | {f.get('gold_id') or ''} {f.get('expected') or ''} | {f.get('got') or ''} | "
                         f"{f.get('expected_value', '')} | {f.get('got_value', '')} | {ev} |")
        lines.append("")
    if planner:
        lines += ["## Planner (member 4's reference cases)", "", f"{planner['passed']}/{planner['total']} cases pass.", "",
                  "| Case | Status (expected) | Scheduled min (expected) | Unscheduled min (expected) | Rule issues |",
                  "| --- | --- | --- | --- | --- |"]
        for r in planner["cases"]:
            lines.append(f"| {r['case']} | {r['status']} ({r['expected_status']}) | {r['scheduled']:g} ({r['expected_scheduled']}) | "
                         f"{r['unscheduled']:g} ({r['expected_unscheduled']}) | {len(r['rule_issues'])} |")
        lines.append("")
    return "\n".join(lines)
